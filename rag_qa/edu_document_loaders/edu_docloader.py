# -*- coding: utf-8 -*-
"""
rag_qa/edu_document_loaders/edu_docloader.py —— Word（.docx）文档加载器
================================================================

【这个文件干什么的？】
   把 Word 文档转成纯文本：段落文字 + 表格文字 + 段落里的图片（OCR）。
   难点在于：Word 的"段落"和"表格"在文档里是交替出现的，
   直接 doc.paragraphs 拿不到表格，直接 doc.tables 又丢了顺序，
   所以本文件用 iter_block_items() 按文档真实顺序逐块遍历。

【技术：python-docx 的两个同名 Document（新手必晕点）】
   from docx import Document as Docu1          # 函数：Docu1(path) 打开文档
   from docx.document import Document as Docu2 # 类：用来判断"对象是不是文档"
   一个用来"打开"，一个用来"做类型判断"，所以导入两次起了别名。

【技术：XML 元素类型判断】
   Word 文档底层是 XML。遍历 doc.element.body 的子元素时：
     CT_P   → 段落（Paragraph）
     CT_Tbl → 表格（Table）
   拿到元素后用 Paragraph(child, parent) / Table(child, parent) 包成对象操作。

【技术：从 Word 里抠图片做 OCR】
   段落的 XML 里图片节点的结构是：
     <pic:pic> ... <a:blip r:embed="rId7"/> ...
   通过 XPath 拿到 rId7 这个"关系编号"，
   再用 doc.part.related_parts["rId7"] 找到图片二进制（_blob），
   用 PIL 打开 → numpy 数组 → 交给 OCR。
"""
from typing import Iterator
from rag_qa.edu_document_loaders.edu_ocr import get_ocr  # OCR 工厂

# 导入必要的模块
from tqdm import tqdm                              # 进度条
from docx.table import _Cell, Table                # 表格相关
from docx.oxml.table import CT_Tbl                 # 表格 XML 结构类型
from docx.oxml.text.paragraph import CT_P          # 段落 XML 结构类型
from docx.text.paragraph import Paragraph          # 段落对象
from docx import Document as Docu1                 # 函数：打开 Word 文档
from docx.document import Document as Docu2        # 类：用于类型判断
from docx.parts.image import ImagePart             # Word 内嵌图片的表示
from PIL import Image                              # 图片处理
from io import BytesIO                             # 把内存字节流当文件用
import numpy as np                                 # 数组（喂给 OCR）
from langchain_core.documents import Document      # LangChain 文档对象
from langchain_core.document_loaders import BaseLoader  # 加载器基类


class OCRDOCLoader(BaseLoader):
    """Word 加载器：解析段落 + 表格 + 图片 OCR。"""

    def __init__(self, filepath: str) -> None:
        """初始化：记住要加载的文件路径。"""
        self.filepath = filepath

    def lazy_load(self) -> Iterator[Document]:
        """惰性加载：把整份 Word 合成一个 Document 输出。"""
        line = self.doc2text(self.filepath)  # 核心解析
        yield Document(page_content=line, metadata={"source": self.filepath})

    def doc2text(self, filepath):
        """核心方法：按文档顺序遍历块，拼出全文。"""

        # 创建 OCR 识别对象
        ocr = get_ocr()

        # 读取 Word 文档
        doc = Docu1(filepath)

        # 最终文本结果
        resp = ""

        # 定义一个生成器：按文档真实顺序产出段落/表格对象
        def iter_block_items(parent):
            # 判断 parent 类型：
            if isinstance(parent, Docu2):
                # 文档对象 → 拿 body（正文 XML 根元素）
                parent_elm = parent.element.body
            elif isinstance(parent, _Cell):
                # 单元格对象 → 拿单元格 XML
                parent_elm = parent._tc
            else:
                raise ValueError("OCRDOCLoader parse fail")  # 都不是直接报错

            # iterchildren()：遍历 XML 元素的所有直接子节点（保序！）
            for child in parent_elm.iterchildren():
                if isinstance(child, CT_P):        # 子节点是段落
                    yield Paragraph(child, parent)  # 包装成 Paragraph 产出
                elif isinstance(child, CT_Tbl):     # 子节点是表格
                    yield Table(child, parent)      # 包装成 Table 产出

        # 创建进度条：总块数 ≈ 段落数 + 表格数
        b_unit = tqdm(total=len(doc.paragraphs) + len(doc.tables),
                      desc="OCRDOCLoader block index: 0")

        # 遍历文档中的所有块（段落和表格，按出现顺序）
        for i, block in enumerate(iter_block_items(doc)):
            b_unit.set_description("OCRDOCLoader  block index: {}".format(i))
            b_unit.refresh()  # 刷新进度条

            # ---------- 情况 1：块是段落 ----------
            if isinstance(block, Paragraph):
                resp += block.text.strip() + "\n"  # 段落文字拼进结果（去首尾空白）

                # 获取段落中的所有图片（XPath 查找 pic:pic 节点）
                images = block._element.xpath('.//pic:pic')
                for image in images:
                    # 从图片节点里取关系编号 r:embed（如 "rId7"）
                    for img_id in image.xpath('.//a:blip/@r:embed'):
                        # 用关系编号在文档的"关系表"里找到对应的图片部件
                        part = doc.part.related_parts[img_id]
                        if isinstance(part, ImagePart):  # 确认是图片
                            # BytesIO 把内存里的二进制数据包装成"类文件对象"
                            # part._blob 是图片的二进制内容
                            image = Image.open(BytesIO(part._blob))
                            # PIL 图片 → numpy 数组 → OCR 识别
                            result, _ = ocr(np.array(image))
                            if result:
                                ocr_result = [line[1] for line in result]  # 取识别文字
                                resp += "\n".join(ocr_result)              # 拼进结果

            # ---------- 情况 2：块是表格 ----------
            elif isinstance(block, Table):
                # 三层循环：行 → 单元格 → 单元格里的段落
                for row in block.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            resp += paragraph.text.strip() + "\n"  # 单元格文字拼进结果

            # 更新进度条
            b_unit.update(1)

        # 返回提取的全部文本
        return resp


# ============================================================
# 【单独调试】python -m rag_qa.edu_document_loaders.edu_docloader
# 注意把下面路径改成你本机的真实 Word 路径再运行。
# ============================================================
if __name__ == '__main__':
    docx_loader = OCRDOCLoader(filepath='/0001.项目目录/others/test/人工智能就业课课程大纲.docx')
    doc = docx_loader.load()
    print(doc)
